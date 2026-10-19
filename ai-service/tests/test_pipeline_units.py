from rag.cleaner import clean_pages, normalize_text
from rag.loader import Page, load_document, DocumentLoadError
from rag.splitter import detect_clause_ref, split_pages
import pytest


def test_normalize_collapses_whitespace_and_keeps_paragraphs():
    out = normalize_text("Hello   world\u00a0\u200b!\n\n\n  Second   para \n continues")
    assert out == "Hello world !\n\nSecond para continues"


def test_cleaner_removes_repeating_headers_and_page_numbers():
    pages = [
        Page(i, f"ACME CONFIDENTIAL\n\nBody text number {i} is long enough to stay in the document.\n\nPage {i} of 4")
        for i in range(1, 5)
    ]
    cleaned = clean_pages(pages)
    for p in cleaned:
        assert "ACME CONFIDENTIAL" not in p.text
        assert "of 4" not in p.text
        assert "Body text number" in p.text


def test_splitter_respects_size_overlap_and_pages():
    para = "The tenant shall pay rent on time. " * 80
    chunks = split_pages([Page(1, para), Page(2, para)], chunk_size=300, overlap=50)
    assert len(chunks) > 4
    assert all(len(c.text) <= 300 for c in chunks)
    assert {c.page for c in chunks} == {1, 2}
    assert [c.chunk_index for c in chunks] == list(range(len(chunks)))


@pytest.mark.parametrize(
    "text,expected",
    [
        ("Clause 8.2 The tenant must give notice.", "Clause 8.2"),
        ("Section 4 Payment terms apply.", "Section 4"),
        ("8.2 Termination. The tenant must give notice.", "Clause 8.2"),
        ("2. Rent and Payment\nThe Tenant shall pay", "Clause 2"),
        ("The tenant must give notice.", None),
    ],
)
def test_clause_reference_detection(text, expected):
    assert detect_clause_ref(text) == expected


def test_txt_loader_and_virtual_pages(tmp_path):
    f = tmp_path / "a.txt"
    f.write_text("\n\n".join(f"Paragraph {i}. " + "word " * 120 for i in range(8)), encoding="utf-8")
    pages = load_document(str(f), "txt")
    assert len(pages) > 1
    assert pages[0].number == 1


def test_docx_loader_reads_paragraphs_and_tables(tmp_path):
    from docx import Document

    d = Document()
    d.add_paragraph("Service Agreement between Acme Ltd and Beta LLC.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Fee"
    t.rows[0].cells[1].text = "Rs. 10,000"
    path = tmp_path / "s.docx"
    d.save(path)
    text = " ".join(p.text for p in load_document(str(path), "docx"))
    assert "Acme Ltd" in text and "Fee | Rs. 10,000" in text


def test_empty_file_is_rejected(tmp_path):
    f = tmp_path / "empty.txt"
    f.write_text("   \n  ")
    with pytest.raises(DocumentLoadError):
        load_document(str(f), "txt")


def test_pdf_loader_reads_real_pages_in_order(tmp_path):
    import fitz

    path = tmp_path / "contract.pdf"
    doc = fitz.open()
    for n in (1, 2, 3):
        page = doc.new_page()
        page.insert_textbox(fitz.Rect(72, 72, 520, 400), f"Clause {n}. The tenant must give thirty days notice on page {n}.", fontsize=11)
    doc.save(path)
    doc.close()

    pages = load_document(str(path), "pdf")
    assert [p.number for p in pages] == [1, 2, 3]
    assert "page 2" in pages[1].text


def test_pdf_without_text_gives_helpful_error(tmp_path):
    import fitz

    path = tmp_path / "scan.pdf"
    doc = fitz.open()
    doc.new_page()
    doc.save(path)
    doc.close()
    with pytest.raises(DocumentLoadError, match="scanned"):
        load_document(str(path), "pdf")


def test_generator_renumbers_citations(monkeypatch):
    from rag import generator
    from rag.retriever import RetrievedChunk

    chunks = [RetrievedChunk(f"passage {i}", i, None, 0.5, i) for i in range(1, 5)]

    class FakeLLM:
        is_mock = False

        def complete(self, system, user, max_tokens=0, temperature=0):
            return "Notice is 30 days [3], and renewal is automatic [1][3]. Ignore [9]."

    monkeypatch.setattr(generator, "get_llm", lambda: FakeLLM())
    answer, found, sources = generator.generate_answer("q", chunks)
    assert found and [s.page for s in sources] == [1, 3]
    assert answer == "Notice is 30 days [2], and renewal is automatic [1][2]. Ignore ."
