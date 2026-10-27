"""FastAPI AI service: document ingestion, retrieval-augmented Q&A, summaries, risk and clause analysis."""
import logging
import os
import re
import secrets
import tempfile
from difflib import SequenceMatcher
from contextlib import asynccontextmanager

from fastapi import Depends, FastAPI, File, Form, Header, HTTPException, Request, UploadFile
from fastapi.responses import JSONResponse, Response

from config import settings
from models.schemas import (
    AskRequest, AskResponse, ClauseResponse, CompareRequest, DocRef, EmbedRequest, ProcessResponse,
    ReportExportRequest, RetrieveRequest, RiskResponse, Source,
)
from rag import vector_store
from rag.embeddings import embed_texts, warm_up
from rag.generator import condense_question, format_skill_list, generate_answer
from rag.loader import DocumentLoadError
from rag.retriever import retrieve
from services import clause_detector, document_service, risk_analyzer, summarizer, text_store
from services.llm_client import LLMError, get_llm

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s: %(message)s")
log = logging.getLogger("ai-service")

ALLOWED_TYPES = {"pdf", "docx", "txt"}
SUMMARY_REQUEST = re.compile(
    r"\b(?:summari[sz](?:e|ation)|summaris(?:e|ation))\b"
    r"|\b(?:explain|give|provide|write|create|make)\s+(?:me\s+)?(?:a\s+)?(?:(?:brief|short|detailed|executive|overall|full|quick)\s+)?(?:summary|overview)\b"
    r"|\b(?:(?:brief|short|detailed|executive|overall|full|quick)\s+)?(?:summary|overview)\s+(?:of|for|about)\s+(?:the\s+|this\s+|my\s+|uploaded\s+)?(?:document|file|resume|contract)\b"
    r"|^\s*(?:(?:brief|short|detailed|executive|overall|full|quick)\s+)?(?:summary|overview)\s*[.!?]*\s*$"
    r"|\b(?:what is|what's)\s+(?:this|the)\s+(?:document|file|resume|contract)\s+about\b"
    r"|\bkey\s+points\s+(?:of|from)\s+(?:the\s+|this\s+|my\s+|uploaded\s+)?(?:document|file|resume|contract)\b",
    re.IGNORECASE,
)
SUMMARY_PAGE_SCOPE = re.compile(r"\bpages?\s+(\d+)(?:\s*(?:-|to|through)\s*(\d+))?\b", re.IGNORECASE)
SKILL_REQUEST = re.compile(r"\bskills?\b|\bcompetenc(?:y|ies)\b", re.IGNORECASE)
EMAIL_QUESTION = re.compile(r"\b(?:e-?mail|mail\s+address)\b", re.IGNORECASE)
EMAIL_ADDRESS = re.compile(r"[a-z0-9.!#$%&'*+/=?^_`{|}~-]+@[a-z0-9](?:[a-z0-9-]{0,61}[a-z0-9])?(?:\.[a-z0-9](?:[a-z0-9-]{0,61}[a-z0-9])?)+", re.IGNORECASE)
DATE_QUESTION = re.compile(r"\b(?:date|dated|when)\b", re.IGNORECASE)
TIME_QUESTION = re.compile(r"\btime\b", re.IGNORECASE)
PRESENCE_QUESTION = re.compile(
    r"\b(?:does|do)\s+(?:it|this|the document|my document)\s+(?:contain|include|mention|have)\b"
    r"|\b(?:is|are)\s+(?:it|this|that|there|the document|my document)\b",
    re.IGNORECASE,
)
MONTHS = r"Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:t(?:ember)?)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?"
DOCUMENT_DATE = re.compile(
    rf"\b(?:\d{{1,2}}\s+(?:{MONTHS})\s+\d{{4}}|(?:{MONTHS})\s+\d{{1,2}},?\s+\d{{4}}|\d{{4}}[-/]\d{{1,2}}[-/]\d{{1,2}}|\d{{1,2}}[-/]\d{{1,2}}[-/]\d{{2,4}})\b",
    re.IGNORECASE,
)
DOCUMENT_TIME = re.compile(r"\b(?:[01]?\d|2[0-3]):[0-5]\d(?:\s?[AP]M)?\b|\b\d{1,2}\s?[AP]M\b", re.IGNORECASE)


def _has_entity_intent(question: str, intent: str, pattern: re.Pattern) -> bool:
    if pattern.search(question):
        return True
    # Accept a one-character typo in short intent words, e.g. "timw" for "time".
    return any(
        len(token) == len(intent) and sum(a != b for a, b in zip(token, intent)) == 1
        for token in re.findall(r"[a-z]+", question.lower())
    )


def _extract_document_values(pages, pattern: re.Pattern, label: str, question: str):
    sources: list[Source] = []
    lines: list[str] = []
    source_by_page: dict[int, int] = {}
    seen: set[str] = set()
    for page in pages:
        for match in pattern.finditer(page.text):
            value = match.group(0).strip(" .,;:")
            normalized = value.casefold()
            if normalized in seen:
                continue
            seen.add(normalized)
            if page.number not in source_by_page:
                start = max(0, match.start() - 160)
                end = min(len(page.text), match.end() + 160)
                source_by_page[page.number] = len(sources) + 1
                sources.append(Source(page=page.number, text=page.text[start:end].strip(), score=1.0))
            lines.append(f"- {value} [{source_by_page[page.number]}]")
    if not lines:
        return None
    return AskResponse(
        answer=f"{label} found in the document:\n" + "\n".join(lines),
        found=True,
        standalone_question=question,
        sources=sources,
    )


def _presence_subject(question: str) -> str:
    subject = re.sub(
        r"^\s*(?:(?:does|do)\s+(?:it|this|the document|my document)\s+(?:contain|include|mention|have)|"
        r"(?:is|are)\s+(?:it|this|that|there|the document|my document))\s+",
        "",
        question,
        flags=re.IGNORECASE,
    )
    return subject.strip(" ?.!\t\r\n") or "that item"


@asynccontextmanager
async def lifespan(app: FastAPI):
    llm = get_llm()
    log.info("LLM provider=%s model=%s | embeddings=%s", llm.provider, settings.llm_model, settings.embedding_model)
    if os.getenv("SKIP_WARMUP") != "1":
        warm_up()  # load the embedding model once at startup instead of on the first request
    yield


app = FastAPI(title="Legal Analyzer AI Service", version="1.0.0", lifespan=lifespan)


def require_key(x_internal_key: str = Header(default="")) -> None:
    """Only the Node backend (which knows the shared secret) may call this service."""
    if not secrets.compare_digest(x_internal_key.encode(), settings.internal_key.encode()):
        raise HTTPException(status_code=401, detail="Invalid internal key")


# ---------------------------------------------------------------- error mapping
@app.exception_handler(DocumentLoadError)
async def _load_error(_: Request, exc: DocumentLoadError):
    return JSONResponse(status_code=422, content={"detail": str(exc)})


@app.exception_handler(text_store.DocumentNotFound)
async def _not_found(_: Request, exc: text_store.DocumentNotFound):
    return JSONResponse(status_code=404, content={"detail": str(exc)})


@app.exception_handler(LLMError)
async def _llm_error(_: Request, exc: LLMError):
    log.error("LLM error: %s", exc)
    return JSONResponse(status_code=502, content={"detail": str(exc)})


@app.exception_handler(ValueError)
async def _value_error(_: Request, exc: ValueError):
    return JSONResponse(status_code=400, content={"detail": str(exc)})


# ---------------------------------------------------------------- routes
@app.get("/health")
def health():
    llm = get_llm()
    return {"status": "ok", "llm_provider": llm.provider, "embedding_model": settings.embedding_model}


@app.post("/process-document", response_model=ProcessResponse, dependencies=[Depends(require_key)])
def process_document(
    file: UploadFile = File(...),
    document_id: str = Form(...),
    user_id: str = Form(...),
    file_type: str = Form(...),
):
    file_type = file_type.lower()
    if file_type not in ALLOWED_TYPES:
        raise HTTPException(400, "Unsupported file type")
    text_store.safe_id(document_id), text_store.safe_id(user_id)

    limit = settings.max_upload_mb * 1024 * 1024
    with tempfile.NamedTemporaryFile(suffix=f".{file_type}", delete=False) as tmp:
        size = 0
        while chunk := file.file.read(1024 * 1024):
            size += len(chunk)
            if size > limit:
                tmp.close()
                os.unlink(tmp.name)
                raise HTTPException(413, f"File is larger than {settings.max_upload_mb} MB")
            tmp.write(chunk)
        tmp_path = tmp.name
    try:
        return document_service.process_document(tmp_path, file_type, user_id, document_id)
    finally:
        os.unlink(tmp_path)


@app.post("/ask", response_model=AskResponse, dependencies=[Depends(require_key)])
def ask(req: AskRequest):
    if vector_store.count(req.user_id, req.document_id) == 0:
        raise text_store.DocumentNotFound("This document has not been processed yet")

    # Exact fact questions should not depend on a semantic similarity score
    # or an LLM. Extract common contact details directly from document text.
    pages = None
    if EMAIL_QUESTION.search(req.question):
        pages = text_store.load_pages(req.user_id, req.document_id)
        answer = _extract_document_values(pages, EMAIL_ADDRESS, "Email address(es)", req.question)
        if answer:
            return answer
    if _has_entity_intent(req.question, "date", DATE_QUESTION):
        pages = pages or text_store.load_pages(req.user_id, req.document_id)
        answer = _extract_document_values(pages, DOCUMENT_DATE, "Date(s)", req.question)
        if answer:
            return answer
    if _has_entity_intent(req.question, "time", TIME_QUESTION):
        pages = pages or text_store.load_pages(req.user_id, req.document_id)
        answer = _extract_document_values(pages, DOCUMENT_TIME, "Time(s)", req.question)
        if answer:
            return answer

    if SKILL_REQUEST.search(req.question):
        pages = text_store.load_pages(req.user_id, req.document_id)
        answer = format_skill_list("\n".join(page.text for page in pages))
        if answer:
            chunks = retrieve(req.question, req.user_id, req.document_id, top_k=req.top_k, min_score=0.0)
            return AskResponse(
                answer=answer,
                found=True,
                standalone_question=req.question,
                sources=[Source(page=c.page, clause=c.clause, text=c.text, score=c.score) for c in chunks],
            )

    if SUMMARY_REQUEST.search(req.question):
        pages = text_store.load_pages(req.user_id, req.document_id)
        scope = SUMMARY_PAGE_SCOPE.search(req.question)
        if scope:
            start_page = int(scope.group(1))
            end_page = int(scope.group(2) or start_page)
            if end_page < start_page:
                start_page, end_page = end_page, start_page
            selected_pages = [p for p in pages if start_page <= p.number <= end_page]
            if not selected_pages:
                last_page = max((p.number for p in pages), default=0)
                answer = f"Page {start_page} is outside this document's page range (1–{last_page})."
                return AskResponse(
                    answer=answer,
                    found=False,
                    standalone_question=req.question,
                    sources=[],
                )
            scope_label = (
                f"page {start_page}" if start_page == end_page
                else f"pages {start_page}–{end_page}"
            )
            answer = summarizer.summarize_for_chat(selected_pages, scope_label=scope_label)
        else:
            answer = summarizer.summarize_for_chat(pages)
        return AskResponse(
            answer=answer or "I could not find readable text in the uploaded document.",
            found=bool(answer),
            standalone_question=req.question,
            sources=[],
        )

    history = [h.model_dump() for h in req.history]
    presence_check = bool(PRESENCE_QUESTION.search(req.question))
    # In "does it contain X?", "it" refers to the selected document, so do not
    # spend another model call rewriting it from chat history.
    standalone = req.question if presence_check else condense_question(req.question, history)
    # For explicit presence checks, let the model inspect the closest passages
    # even when their similarity is below the normal answer threshold.
    # This allows useful answers such as distinguishing an assessment invite
    # from an offer letter, while still grounding the response in cited text.
    min_score = 0.0 if presence_check else None
    chunks = retrieve(standalone, req.user_id, req.document_id, top_k=req.top_k, min_score=min_score)
    answer, found, used = generate_answer(standalone, chunks)
    if presence_check and not found and chunks:
        # If the small local model still declines to answer, show the closest
        # passage and state the limit of what it supports instead of a dead end.
        answer = (
            f'I don\'t see "{_presence_subject(standalone)}" identified in the passages I found. '
            f'The closest passage says: "{chunks[0].text[:500].strip()}" [1]'
        )
        found, used = True, chunks[:1]

    return AskResponse(
        answer=answer,
        found=found,
        standalone_question=standalone,
        sources=[Source(page=c.page, clause=c.clause, text=c.text, score=c.score) for c in used],
    )


@app.post("/retrieve", dependencies=[Depends(require_key)])
def retrieve_endpoint(req: RetrieveRequest):
    """Debug/inspection endpoint: shows what the retriever returns for a query."""
    chunks = retrieve(req.query, req.user_id, req.document_id, top_k=req.top_k, min_score=0.0)
    return {"chunks": [c.__dict__ for c in chunks]}


@app.post("/embed", dependencies=[Depends(require_key)])
def embed(req: EmbedRequest):
    vectors = embed_texts(req.texts)
    return {"model": settings.embedding_model, "dimensions": len(vectors[0]), "embeddings": vectors}


@app.post("/summarize", dependencies=[Depends(require_key)])
def summarize(req: DocRef):
    return summarizer.summarize(text_store.load_pages(req.user_id, req.document_id))


@app.post("/analyze-risk", response_model=RiskResponse, dependencies=[Depends(require_key)])
def analyze_risk(req: DocRef):
    return risk_analyzer.analyze(text_store.load_pages(req.user_id, req.document_id))


@app.post("/detect-clauses", response_model=ClauseResponse, dependencies=[Depends(require_key)])
def detect_clauses(req: DocRef):
    return clause_detector.detect(text_store.load_pages(req.user_id, req.document_id))


@app.post("/compare-documents", dependencies=[Depends(require_key)])
def compare_documents(req: CompareRequest):
    if req.document_id == req.other_document_id:
        raise HTTPException(400, "Choose two different documents")
    old_pages = text_store.load_pages(req.user_id, req.document_id)
    new_pages = text_store.load_pages(req.user_id, req.other_document_id)
    old_text = "\n".join(page.text for page in old_pages)
    new_text = "\n".join(page.text for page in new_pages)
    if max(len(old_text), len(new_text)) > 500000:
        raise HTTPException(413, "Documents are too large to compare in one request")

    old_lines = [line.strip() for line in old_text.splitlines() if line.strip()]
    new_lines = [line.strip() for line in new_text.splitlines() if line.strip()]
    opcodes = [
        opcode for opcode in SequenceMatcher(a=old_lines, b=new_lines).get_opcodes()
        if opcode[0] != "equal"
    ]
    changes = []
    for tag, i1, i2, j1, j2 in opcodes[:100]:
        changes.append({
            "type": tag,
            "old_text": "\n".join(old_lines[i1:i2])[:3000],
            "new_text": "\n".join(new_lines[j1:j2])[:3000],
        })

    old_risk = risk_analyzer.analyze(old_pages)
    new_risk = risk_analyzer.analyze(new_pages)
    old_by_category = {risk["category"]: risk for risk in old_risk["risks"]}
    new_by_category = {risk["category"]: risk for risk in new_risk["risks"]}
    finding_details = lambda findings: [
        {
            "category": category,
            "title": finding["title"],
            "level": finding["level"],
            "explanation": finding["explanation"],
        }
        for category, finding in findings.items()
    ]
    return {
        "changes": changes,
        "risk_change": {
            "before": old_risk["risk_level"],
            "after": new_risk["risk_level"],
            "added_findings": finding_details({k: v for k, v in new_by_category.items() if k not in old_by_category}),
            "removed_findings": finding_details({k: v for k, v in old_by_category.items() if k not in new_by_category}),
        },
        "truncated": len(opcodes) > 100,
    }


@app.post("/export-report", dependencies=[Depends(require_key)])
def export_report(req: ReportExportRequest):
    import io
    from html import escape

    if req.format == "docx":
        from docx import Document

        output = io.BytesIO()
        document = Document()
        for line in req.report_text.splitlines():
            if line.startswith("# "):
                document.add_heading(line[2:], level=1)
            elif line.startswith("## "):
                document.add_heading(line[3:], level=2)
            elif line.startswith("### "):
                document.add_heading(line[4:], level=3)
            elif line.startswith("- "):
                document.add_paragraph(line[2:], style="List Bullet")
            elif line.startswith("> "):
                document.add_paragraph(line[2:], style="Quote")
            elif line.strip():
                document.add_paragraph(line.replace("**", ""))
        document.save(output)
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

        output = io.BytesIO()
        styles = getSampleStyleSheet()
        body = ParagraphStyle("Body", parent=styles["BodyText"], leading=14, spaceAfter=7)
        story = []
        for line in req.report_text.splitlines():
            if not line.strip():
                story.append(Spacer(1, 5))
                continue
            level = 0
            while level < len(line) and line[level] == "#":
                level += 1
            if level:
                style = styles["Title"] if level == 1 else styles["Heading2"] if level == 2 else styles["Heading3"]
                content = line[level:].strip()
                story.append(Paragraph(escape(content), style))
            else:
                content = line.strip()
                is_bullet = content.startswith("- ")
                if content.startswith("- "):
                    content = content[2:]
                elif content.startswith("> "):
                    content = content[2:]
                content = escape(content).replace("**", "")
                if is_bullet:
                    content = "&#8226; " + content
                story.append(Paragraph(content, body))
        SimpleDocTemplate(output, pagesize=letter, rightMargin=0.7 * inch, leftMargin=0.7 * inch,
                          topMargin=0.7 * inch, bottomMargin=0.7 * inch).build(story)
        media_type = "application/pdf"
    return Response(output.getvalue(), media_type=media_type)


@app.delete("/documents/{document_id}", dependencies=[Depends(require_key)])
def delete_document(document_id: str, user_id: str):
    document_service.delete_document(user_id, document_id)
    return {"deleted": True}
